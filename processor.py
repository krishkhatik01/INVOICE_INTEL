import os
import json
import tempfile
import pandas as pd
import fitz  # PyMuPDF


class DocumentProcessor:
    """Processes mixed document formats locally without external APIs.

    Returns either:
      - plain text content (for text/tabular documents)
      - "[IMAGE_DATA]:<path>" for a single image to be sent to a vision model
      - "[ERROR PARSING DOCUMENT]: <reason>" on failure
    """

    # Extensions we know how to handle locally
    SUPPORTED_EXTENSIONS = {
        ".pdf", ".txt", ".md", ".markdown", ".csv", ".xlsx", ".xls",
        ".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp",
        ".docx", ".eml", ".json",
    }

    @staticmethod
    def process_file(file_path: str) -> str:
        ext = os.path.splitext(file_path)[1].lower()

        try:
            if ext == ".pdf":
                return DocumentProcessor._process_pdf(file_path)

            elif ext in [".txt", ".md", ".markdown"]:
                with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                    return f.read()

            elif ext == ".csv":
                df = pd.read_csv(file_path)
                return df.to_string()

            elif ext in [".xlsx", ".xls"]:
                # Read every sheet so multi-sheet workbooks aren't silently truncated
                sheets = pd.read_excel(file_path, sheet_name=None)
                parts = []
                for sheet_name, df in sheets.items():
                    parts.append(f"--- Sheet: {sheet_name} ---\n{df.to_string()}")
                return "\n\n".join(parts)

            elif ext in [".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp"]:
                return f"[IMAGE_DATA]: {file_path}"

            elif ext == ".docx":
                return DocumentProcessor._process_docx(file_path)

            elif ext == ".eml":
                return DocumentProcessor._process_eml(file_path)

            elif ext == ".json":
                with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                    data = json.load(f)
                return json.dumps(data, indent=2)

            else:
                return f"Unsupported file type: {ext}"

        except Exception as e:
            return f"[ERROR PARSING DOCUMENT]: {str(e)}"

    # ------------------------------------------------------------------
    # Format-specific helpers
    # ------------------------------------------------------------------

    @staticmethod
    def _process_pdf(file_path: str) -> str:
        doc = fitz.open(file_path)
        text = "\n".join(page.get_text("text") for page in doc)

        # If the PDF has (almost) no extractable text, it's likely a scanned
        # document/image-based invoice. Render the first page as an image so
        # it can be routed to the vision model instead of being skipped.
        if len(text.strip()) < 20 and len(doc) > 0:
            page = doc[0]
            pix = page.get_pixmap(matrix=fitz.Matrix(2, 2))  # ~144 DPI
            tmp_dir = tempfile.gettempdir()
            base_name = os.path.splitext(os.path.basename(file_path))[0]
            image_path = os.path.join(tmp_dir, f"{base_name}_page1_render.png")
            pix.save(image_path)
            doc.close()
            return f"[IMAGE_DATA]: {image_path}"

        doc.close()
        return text

    @staticmethod
    def _process_docx(file_path: str) -> str:
        try:
            from docx import Document
        except ImportError:
            return "[ERROR PARSING DOCUMENT]: python-docx is not installed"

        doc = Document(file_path)
        parts = []

        for para in doc.paragraphs:
            if para.text.strip():
                parts.append(para.text)

        for table in doc.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                parts.append(" | ".join(cells))

        return "\n".join(parts)

    @staticmethod
    def _process_eml(file_path: str) -> str:
        from email import policy
        from email.parser import BytesParser

        with open(file_path, "rb") as f:
            msg = BytesParser(policy=policy.default).parse(f)

        headers = (
            f"From: {msg.get('From', 'Unknown')}\n"
            f"To: {msg.get('To', 'Unknown')}\n"
            f"Subject: {msg.get('Subject', 'Unknown')}\n"
            f"Date: {msg.get('Date', 'Unknown')}\n"
        )

        body = ""
        if msg.is_multipart():
            for part in msg.walk():
                if part.get_content_type() == "text/plain" and not part.get_filename():
                    body += part.get_content()
        else:
            if msg.get_content_type() == "text/plain":
                body = msg.get_content()

        attachments = [part.get_filename() for part in msg.iter_attachments() if part.get_filename()]
        attachment_note = f"\nAttachments: {', '.join(attachments)}" if attachments else ""

        return f"{headers}{attachment_note}\n\n{body}"
